from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any


def _infer_scalar(values: list[str]) -> str:
    non_empty = [value for value in values if value not in ("", None)]
    if not non_empty:
        return "string"
    try:
        for value in non_empty:
            int(value)
        return "integer"
    except (ValueError, TypeError):
        pass
    try:
        for value in non_empty:
            float(value)
        return "number"
    except (ValueError, TypeError):
        pass
    lowered = {str(value).lower() for value in non_empty}
    if lowered <= {"true", "false", "yes", "no", "0", "1"}:
        return "boolean"
    return "string"


def _tabular_profile(rows: list[dict[str, Any]]) -> dict[str, Any]:
    columns = list(rows[0].keys()) if rows else []
    column_profiles = []
    for column in columns:
        values = ["" if row.get(column) is None else str(row.get(column)) for row in rows]
        filled = [value for value in values if value != ""]
        top_values = Counter(filled).most_common(3)
        column_profiles.append(
            {
                "name": column,
                "inferred_type": _infer_scalar(values),
                "null_count": len(values) - len(filled),
                "distinct_count": len(set(filled)),
                "top_values": [{"value": value, "count": count} for value, count in top_values],
            }
        )
    return {
        "kind": "structured",
        "row_count": len(rows),
        "column_count": len(columns),
        "columns": column_profiles,
        "sample_rows": rows[:8],
    }


def read_structured_rows(path: Path, suffix: str) -> list[dict[str, Any]] | None:
    suffix = suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list) and (not payload or isinstance(payload[0], dict)):
            return payload
        return None
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(values[0])]
        return [dict(zip(headers, row, strict=False)) for row in values[1:]]
    if suffix == ".parquet":
        import pyarrow.parquet as parquet

        table = parquet.read_table(path)
        return table.to_pylist()
    return None


def profile_file(path: Path, suffix: str) -> dict[str, Any]:
    rows = read_structured_rows(path, suffix)
    if rows is not None:
        return _tabular_profile(rows)
    suffix = suffix.lower()
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        return {
            "kind": "document",
            "row_count": None,
            "column_count": None,
            "preview": json.dumps(payload, indent=2)[:4000],
        }
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return {
            "kind": "document",
            "page_count": len(reader.pages),
            "character_count": len(text),
            "preview": text[:4000],
            "metadata": {str(key): str(value) for key, value in (reader.metadata or {}).items()},
        }
    return {
        "kind": "binary",
        "message": "Metadata stored. Content profiling is not available for this file type.",
    }
