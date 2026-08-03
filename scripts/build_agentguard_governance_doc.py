from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "docs" / "AgentGuard_Governance_Integration_Guide.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F4F6F9"
RISK = "9B1C1C"
USABLE_WIDTH = 6.5


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(round(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths, strict=False):
        grid_col.set(qn("w:w"), str(round(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=False):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def write_text(paragraph, text: str, **kwargs) -> None:
    run = paragraph.add_run(text)
    set_font(run, **kwargs)


def set_style(style, size: float, color: str, before: float, after: float, bold: bool = False) -> None:
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.15


def add_heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    write_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)


def add_body(doc, text: str, after: float = 8) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(after)
    write_text(p, text, size=10.5, color="1F2937")


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        write_text(p, item, size=10.5, color="1F2937")


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        write_text(p, item, size=10.5, color="1F2937")


def add_callout(doc, label: str, text: str, color: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    write_text(p, label + " ", size=10.5, color=INK, bold=True)
    write_text(p, text, size=10.5, color="1F2937")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers, strict=False):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        write_text(p, header, size=9.5, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=False):
            p = cell.paragraphs[0]
            write_text(p, value, size=9.0, color="1F2937")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_source(doc, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(3)
    write_text(p, text, size=8.5, color=MUTED)


def make_document() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    normal = doc.styles["Normal"]
    set_style(normal, 10.5, "1F2937", 0, 8)
    set_style(doc.styles["Heading 1"], 16, BLUE, 14, 7, True)
    set_style(doc.styles["Heading 2"], 13, BLUE, 12, 6, True)
    set_style(doc.styles["Heading 3"], 12, DARK_BLUE, 10, 5, True)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    write_text(header, "DataPilot Agent OS | Governance Integration Guide", size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    write_text(footer, "Internal technical reference | ", size=8.5, color=MUTED)
    write_text(footer, str(date.today()), size=8.5, color=MUTED)

    # Memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    write_text(p, "GOVERNANCE ARCHITECTURE", size=9.5, color=DARK_BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    write_text(p, "AgentGuard Integration and Provider-Neutral Governance", size=23, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    write_text(p, "Purpose, implemented coverage, market comparison, gaps, and operating model", size=12, color=MUTED)
    add_matrix(doc, ["Document owner", "Scope", "Decision"], [["Engineering", "DataPilot API and worker", "Use a generic governance contract with configurable adapter fan-out"]], [1.4, 2.4, 2.7])

    add_callout(doc, "Executive decision.", "AgentGuard is useful beyond LLM calls because its SDK supports custom tracking spans and context. DataPilot now uses a provider-neutral contract that can fan out the same sanitized events to AgentGuard, an OTLP/HTTP collector (including Phoenix or Langfuse), and a generic signed HTTPS webhook. None of these adapters replace the application's local authorization, approvals, or audit enforcement.")

    add_heading(doc, "1. What AgentGuard does in this architecture")
    add_body(doc, "AgentGuard is treated as a telemetry destination. It receives structured, sanitized observations so operators can see the execution path, group it by project/user/session/feature, and analyze model usage. The application keeps the ability to allow, deny, approve, and audit actions locally, which avoids placing runtime authorization availability on an external observability service.")
    add_matrix(doc, ["Layer", "Responsibility", "System of record"], [
        ["Policy enforcement", "Authorization, read-only constraints, allowlists, human approvals", "DataPilot application and database"],
        ["Governance telemetry", "Execution metadata, outcome, duration, risk classification, LLM usage", "Selected provider adapter (AgentGuard today)"],
        ["Evidence and audit", "Approval evidence, durable job logs, local audit events", "DataPilot database"],
    ], [1.4, 3.1, 2.0])

    add_heading(doc, "2. Implemented feature utilization")
    add_matrix(doc, ["Feature", "Implementation", "Data sent"], [
        ["Initialization", "API startup and Temporal worker startup initialize the configured adapter.", "Service name, environment; no prompt data."],
        ["Model generation", "OpenAI-compatible, Gemini, Claude, and company HTTP calls emit a generation record.", "Model, input/output token use, feature, outcome, project/user/session."],
        ["Agent and feature attribution", "Planner, SQL generation, SQL repair, diagnosis, and evaluations are labelled as features.", "Stable feature names and project/user/run identifiers."],
        ["Tool execution", "Built-in and allowlisted HTTP tools emit success/failure, retry count, and duration.", "Tool handler, implementation type, attempts, duration; never parameters."],
        ["Connector query", "External read-only connector queries emit a success observation.", "Connector type/id, read-only flag, row count, duration; never SQL or credentials."],
        ["Approval decision", "Approval decisions emit the action type, risk level, and approved/rejected outcome.", "Approval id, action type, risk, actor, job/session."],
        ["Privacy controls", "Content capture and infrastructure spans remain disabled by default.", "Metadata only unless configuration is deliberately changed."],
    ], [1.35, 3.1, 2.05])
    add_callout(doc, "Important distinction.", "AgentGuard's documented Python SDK describes observability primitives: initialization, context, feature tracking, and manual generation records. This integration uses its custom spans for operational events. There is no claim that AgentGuard itself blocks a tool call or decides an approval.", PALE_GRAY)

    add_heading(doc, "3. Generic governance contract")
    add_body(doc, "The application no longer calls a vendor from its control points. It constructs a provider-neutral GovernanceEvent or GovernanceGeneration and fans it out to every configured adapter. This allows another vendor, an OpenTelemetry collector, an internal audit stream, or a future policy platform to be added without rewriting model, tool, connector, or approval workflows.")
    add_matrix(doc, ["Application event", "Generic fields", "Adapter delivery"], [
        ["model_generation", "feature, model, token use, outcome, project, user, session", "AgentGuard generation; OTLP GenAI span; webhook generation envelope."],
        ["tool_execution", "outcome, duration, attempts, implementation type", "AgentGuard custom span; OTLP span; webhook event envelope."],
        ["connector_query", "outcome, read-only flag, row count, duration", "AgentGuard custom span; OTLP span; webhook event envelope."],
        ["approval_decision", "decision, risk, action type, approval id", "AgentGuard custom span; OTLP span; webhook event envelope."],
    ], [1.45, 2.9, 2.15])
    add_body(doc, "Privacy boundary: metadata is restricted to primitive values and key names containing credential, key, password, secret, token, or authorization are dropped before an adapter receives the event. SQL text, tool parameters, prompt text, model responses, connector credentials, and approval notes are not exported by this generic layer.")

    add_heading(doc, "4. Minimum configuration and operations")
    add_matrix(doc, ["Variable", "Required", "Meaning"], [
        ["GOVERNANCE_ENABLED", "Yes", "Turns generic telemetry on or off."],
        ["GOVERNANCE_PROVIDERS", "Yes", "Comma-separated fan-out: agentguard, otlp, webhook, or none."],
        ["GOVERNANCE_SERVICE_NAME", "No", "Stable service identity attached to OTLP and webhook events."],
        ["AGENTGUARD_BASE_URL", "For AgentGuard", "AgentGuard OTLP/API endpoint."],
        ["AGENTGUARD_PUBLIC_KEY / SECRET_KEY", "For AgentGuard", "Project-scoped AgentGuard credentials; keep only in a secret store or ignored local env file."],
        ["GOVERNANCE_OTLP_ENDPOINT", "For OTLP", "Vendor-provided complete OTLP/HTTP trace endpoint. Phoenix and Langfuse document OTLP ingestion."],
        ["GOVERNANCE_OTLP_HEADERS_JSON", "For OTLP auth", "JSON header map injected through the deployment secret store; never commit API keys."],
        ["GOVERNANCE_WEBHOOK_URL", "For webhook", "HTTPS endpoint for an internal or vendor governance gateway."],
        ["GOVERNANCE_WEBHOOK_HEADERS_JSON / SIGNING_SECRET", "For webhook auth", "Optional credential headers and HMAC SHA-256 payload signing secret."],
        ["AGENTGUARD_CAPTURE_CONTENT", "No", "Default false. Do not enable without an approved data-handling decision."],
        ["AGENTGUARD_INCLUDE_INFRA_SPANS", "No", "Default false to avoid noisy/expensive infrastructure tracing."],
    ], [2.25, 1.1, 3.15])
    add_numbered(doc, [
        "Set GOVERNANCE_ENABLED=true and choose one or more values in GOVERNANCE_PROVIDERS.",
        "For AgentGuard, set its endpoint and project keys through a secret manager or the ignored local .env file.",
        "For Phoenix, Langfuse, or another collector, set GOVERNANCE_PROVIDERS=otlp and supply its complete OTLP/HTTP trace endpoint and authentication headers as JSON through the secret manager.",
        "For an unsupported vendor, set GOVERNANCE_PROVIDERS=webhook and configure a receiver that accepts the versioned JSON envelope and validates the optional HMAC signature.",
        "Build and restart the API and worker so both processes initialize the adapter.",
        "Generate a SQL request, an agent run, a governed tool call, and an approval decision; verify the four event families in AgentGuard.",
        "Review the telemetry retention, access, and content-capture settings with the data-governance owner before enabling prompt content." 
    ])

    add_heading(doc, "5. Market comparison and documented gaps")
    add_body(doc, "This comparison is intentionally conservative. It contrasts features explicitly documented by AgentGuard's public PyPI package with capabilities documented by representative observability platforms. An item marked 'not documented' is not proof that a capability does not exist in a private console or future AgentGuard release; it requires vendor confirmation before a procurement or compliance conclusion.")
    add_matrix(doc, ["Capability", "AgentGuard SDK evidence", "Comparable market capability / implication"], [
        ["Tracing and LLM cost usage", "Documented: initialization, contexts, feature tags, auto/manual generation tracing.", "Baseline parity for basic observability. Langfuse and Phoenix also expose trace-based analysis."],
        ["Prompt management", "Not documented in the PyPI SDK page.", "Langfuse and Phoenix document prompt versioning, prompt deployment/playgrounds, and replay. Keep DataPilot's versioned agent instructions as the current control."],
        ["Evaluations and experiments", "Not documented in the PyPI SDK page.", "Phoenix and Langfuse document evaluators, datasets, and experiments. DataPilot has local evaluation runs, but no provider-linked score export yet."],
        ["Broad ecosystem instrumentation", "Documented Python SDK and provider auto-instrumentation.", "Phoenix documents OpenTelemetry/OpenInference and broad provider/framework/language integrations. Validate non-Python needs before standardizing."],
        ["Analytics/export ecosystem", "Not documented in the PyPI SDK page.", "Langfuse documents APIs, exports, metrics, CLI and MCP. Confirm AgentGuard's dashboard/export/compliance reporting separately."],
        ["Runtime enforcement", "SDK documentation describes observability, not a policy decision point.", "A production governance program still needs locally enforced authorization, approval, budgets, data controls, and immutable audit evidence."],
    ], [1.45, 2.25, 2.8])

    add_heading(doc, "6. Remaining gaps and recommended next steps")
    add_matrix(doc, ["Priority", "Gap", "Recommended control"], [
        ["High", "No external policy decision point or policy-as-code bundle.", "Add a local policy interface that evaluates pre-execution actions and can fail closed for designated high-risk operations."],
        ["High", "Connector-query failures are not yet exported as a dedicated event.", "Wrap the connector query path in a generic operation helper that emits both success and failure outcomes."],
        ["High", "No pre-call budget or quota enforcement.", "Add project/user/model budgets before provider invocation; export decision and remaining budget as metadata."],
        ["Medium", "No automated PII/secret inspection of prompts and tool arguments.", "Introduce a configurable DLP/secret scanner before provider and external-tool boundaries; default to redaction or approval."],
        ["Medium", "No evaluation-score export to the governance adapter.", "Emit evaluation scores, dataset version, and model/prompt version after local evaluation runs."],
        ["Medium", "No immutable/WORM audit archive.", "Periodically export signed audit evidence to a controlled immutable store if compliance requires it."],
        ["Low", "Vendor-specific enrichment is intentionally not embedded in the application.", "Use OTLP for standard traces; add a dedicated adapter only when a selected vendor requires proprietary features such as prompt deployment or evaluations."],
    ], [0.75, 2.65, 3.1])
    add_callout(doc, "Recommended posture.", "Keep the current minimum integration: generic, metadata-only, fail-open telemetry; local fail-closed controls for high-risk actions. Add adapters and richer capture only after confirming privacy, retention, residency, and enforcement requirements.", PALE_GRAY)

    add_heading(doc, "7. Validation checklist")
    add_bullets(doc, [
        "Confirm model generation events appear for SQL generation and agent planning with model/token metadata.",
        "Confirm a tool execution emits a tool_execution event without parameters or result content.",
        "Confirm an allowlisted connector query emits connector_query with row count and duration only.",
        "Approve or reject an approval request and confirm approval_decision includes risk level and actor attribution.",
        "Inspect a sample event to ensure there are no prompts, SQL, tool arguments, credentials, secrets, or authorization headers.",
        "Test the service with the governance endpoint unavailable; the governed operation must continue and retain its local audit record.",
    ])

    add_heading(doc, "8. Sources and assumptions")
    add_source(doc, "AgentGuard SDK for Python, PyPI project description: https://pypi.org/project/actaclad-agentguard/ (accessed 2026-08-01). Documents initialization, context, feature tracking, manual generation records, span filtering, and Python support.")
    add_source(doc, "Arize Phoenix documentation: https://arize.com/docs/phoenix (accessed 2026-08-01). Documents tracing, evaluation, prompt management, datasets, experiments, OpenTelemetry/OpenInference, and integrations.")
    add_source(doc, "Phoenix OTLP endpoint guidance: https://arize.com/docs/phoenix/learn/faqs/what-is-my-phoenix-endpoint and https://arize.com/docs/phoenix/deployment/authentication (accessed 2026-08-01). Documents trace endpoints and bearer authentication for OTLP exports.")
    add_source(doc, "Langfuse documentation: https://langfuse.com/docs/metrics/overview and https://langfuse.com/docs/api-and-data-platform/overview (accessed 2026-08-01). Documents metrics dimensions, score-based quality, APIs, exports, CLI, and MCP server.")
    add_source(doc, "Langfuse OTLP ingestion: https://langfuse.com/docs/api-and-data-platform/features/public-api and https://langfuse.com/docs/observability/sdk/overview (accessed 2026-08-01). Documents OTLP/HTTP trace ingestion and OpenTelemetry support.")
    add_source(doc, "OpenTelemetry OTLP exporter specification: https://opentelemetry.io/docs/specs/otel/protocol/exporter/ (accessed 2026-08-01). Documents OTLP/HTTP endpoints and header configuration.")
    add_source(doc, "Langfuse experiments documentation: https://langfuse.com/docs/evaluation/experiments/experiments-via-ui (accessed 2026-08-01). Documents dataset-backed prompt/model experiments and evaluators.")
    add_source(doc, "Assumption boundary: market comparison uses public documentation only. Claims of 'not documented' mean unavailable in the reviewed public SDK documentation, not a confirmed product absence.")

    doc.core_properties.title = "AgentGuard Governance Integration Guide"
    doc.core_properties.subject = "DataPilot governance architecture and AgentGuard integration"
    doc.core_properties.author = "DataPilot Engineering"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_document()
